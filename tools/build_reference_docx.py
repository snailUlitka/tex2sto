"""Generate the versioned DOCX reference used by the Pandoc renderer."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import (
    WD_LINE_SPACING,
    WD_PARAGRAPH_ALIGNMENT,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "src" / "tex2sto" / "profiles" / "ssau" / "pandoc" / "reference.docx"
FONT = "Times New Roman"


def _font(style, size: int, *, bold: bool = False, italic: bool = False) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)


def _paragraph(style, *, alignment=None, first_line: bool = True, before=0, after=0) -> None:
    paragraph = style.paragraph_format
    paragraph.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.space_before = Pt(before)
    paragraph.space_after = Pt(after)
    paragraph.first_line_indent = Mm(12.5) if first_line else Mm(0)
    if alignment is not None:
        paragraph.alignment = alignment


def _style(document, name: str, *, size=14, bold=False, italic=False, base="Normal", **kwargs):
    styles = document.styles
    style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = styles[base]
    _font(style, size, bold=bold, italic=italic)
    _paragraph(style, **kwargs)
    return style


def _outline_level(style, level: int) -> None:
    properties = style.element.get_or_add_pPr()
    outline = properties.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        properties.append(outline)
    outline.set(qn("w:val"), str(level))


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, text, end))


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    _font(normal, 14)
    _paragraph(normal, alignment=WD_PARAGRAPH_ALIGNMENT.JUSTIFY)
    _style(document, "Tex2Sto Body", alignment=WD_PARAGRAPH_ALIGNMENT.JUSTIFY)
    _style(
        document,
        "Tex2Sto Structural Heading",
        first_line=False,
        alignment=WD_PARAGRAPH_ALIGNMENT.CENTER,
        before=12,
        after=12,
    )
    for level in range(1, 5):
        style = _style(
            document,
            f"Tex2Sto Heading {level}",
            first_line=False,
            before=12,
            after=6,
        )
        style.paragraph_format.keep_with_next = True
        _outline_level(style, level - 1)
    compact_styles = (
        ("Tex2Sto Figure Content", 14, WD_PARAGRAPH_ALIGNMENT.CENTER),
        ("Tex2Sto Figure Caption", 14, WD_PARAGRAPH_ALIGNMENT.CENTER),
        ("Tex2Sto Table Caption", 14, WD_PARAGRAPH_ALIGNMENT.LEFT),
        ("Tex2Sto Table Text", 12, WD_PARAGRAPH_ALIGNMENT.LEFT),
        ("Tex2Sto Code", 12, WD_PARAGRAPH_ALIGNMENT.LEFT),
        ("Tex2Sto Symbols", 14, WD_PARAGRAPH_ALIGNMENT.LEFT),
        ("Tex2Sto Abstract Keywords", 14, WD_PARAGRAPH_ALIGNMENT.JUSTIFY),
        ("Tex2Sto Title Institution", 12, WD_PARAGRAPH_ALIGNMENT.CENTER),
        ("Tex2Sto Title Details", 14, WD_PARAGRAPH_ALIGNMENT.LEFT),
    )
    for name, size, alignment in compact_styles:
        _style(document, name, size=size, first_line=False, alignment=alignment)
    continuation_style = _style(
        document,
        "Tex2Sto Table Continuation",
        size=12,
        first_line=False,
        alignment=WD_PARAGRAPH_ALIGNMENT.LEFT,
    )
    continuation_style.font.color.rgb = RGBColor(255, 255, 255)
    equation_style = _style(
        document,
        "Tex2Sto Equation",
        first_line=False,
        alignment=WD_PARAGRAPH_ALIGNMENT.LEFT,
    )
    equation_style.paragraph_format.tab_stops.add_tab_stop(
        Mm(82.5),
        WD_TAB_ALIGNMENT.CENTER,
    )
    equation_style.paragraph_format.tab_stops.add_tab_stop(
        Mm(165),
        WD_TAB_ALIGNMENT.RIGHT,
    )
    for level in range(1, 5):
        toc_style = _style(
            document,
            f"Tex2Sto TOC {level}",
            first_line=False,
            alignment=WD_PARAGRAPH_ALIGNMENT.LEFT,
        )
        toc_style.paragraph_format.left_indent = Mm(5 * (level - 1))
        toc_style.paragraph_format.tab_stops.add_tab_stop(
            Mm(165),
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.DOTS,
        )
    table_style = document.styles.add_style("Tex2Sto Table", WD_STYLE_TYPE.TABLE)
    table_style.base_style = document.styles["Table Grid"]
    _font(table_style, 12)
    _style(
        document,
        "Tex2Sto Title Kind",
        bold=True,
        first_line=False,
        alignment=WD_PARAGRAPH_ALIGNMENT.CENTER,
        before=48,
        after=18,
    )
    _style(
        document,
        "Tex2Sto Title Name",
        bold=True,
        first_line=False,
        alignment=WD_PARAGRAPH_ALIGNMENT.CENTER,
        after=36,
    )
    _style(
        document,
        "Tex2Sto Title City",
        first_line=False,
        alignment=WD_PARAGRAPH_ALIGNMENT.CENTER,
        before=72,
    )

    for level in range(1, 5):
        heading = document.styles[f"Heading {level}"]
        _font(heading, 14)

    document.styles["Tex2Sto Table Caption"].paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    _field(footer, "PAGE")
    footer.runs[0].font.name = FONT
    footer.runs[0].font.size = Pt(14)

    document.add_paragraph("tex2sto reference document", style="Tex2Sto Body")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


if __name__ == "__main__":
    main()
