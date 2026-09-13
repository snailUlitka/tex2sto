"""Pandoc-based editable DOCX renderer and narrow OOXML post-processing."""

from __future__ import annotations

import tempfile
import tomllib
from copy import deepcopy
from importlib.resources import files
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, RGBColor

from tex2sto.dialect.syntax import find_command_calls, find_environment
from tex2sto.model import NumberingIndex, ObjectKind, SourceProject
from tex2sto.renderers.docx_fields import (
    add_bookmark,
    add_field,
    repeat_table_row,
    set_update_fields,
)
from tex2sto.renderers.tools import require_tool, require_version, run_tool
from tex2sto.transform import (
    PAGE_BREAK_MARKER,
    SYMBOLS_MARKER,
    TABLE_BREAK_MARKER,
    renderer_body,
)

PROFILE = files("tex2sto") / "profiles" / "ssau" / "pandoc"

STRUCTURAL_TOC_TITLES = {
    "ВВЕДЕНИЕ": "Введение",
    "ЗАКЛЮЧЕНИЕ": "Заключение",
    "ОПРЕДЕЛЕНИЯ, ОБОЗНАЧЕНИЯ И СОКРАЩЕНИЯ": (
        "Определения, обозначения и сокращения"
    ),
    "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ": "Список использованных источников",
}


def _pandoc() -> str:
    executable = require_tool("pandoc")
    configuration = tomllib.loads((PROFILE.parent / "toolchain.toml").read_text())
    require_version(
        [executable, "--version"],
        f"pandoc {configuration['pandoc']}",
        label="Pandoc",
    )
    return executable


def _move_before(element, anchor) -> None:
    anchor.addprevious(element)


def _front_paragraph(document, text: str = "", *, style: str = "Tex2Sto Body"):
    paragraph = document.add_paragraph(text, style=style)
    return paragraph


def _title_page(document, project: SourceProject) -> list:
    metadata = project.metadata
    elements = []
    for text in (
        "МИНОБРНАУКИ РОССИИ",
        "Федеральное государственное автономное образовательное учреждение",
        "высшего образования",
        "«Самарский национальный исследовательский университет",
        "имени академика С. П. Королева»",
        metadata.institute,
        metadata.faculty,
        metadata.department,
    ):
        if not text:
            continue
        paragraph = _front_paragraph(document, text, style="Tex2Sto Title Institution")
        elements.append(paragraph._p)
    title_kind = _front_paragraph(
        document,
        "ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА",
        style="Tex2Sto Title Kind",
    )
    elements.append(title_kind._p)
    title = _front_paragraph(document, metadata.title.upper(), style="Tex2Sto Title Name")
    elements.append(title._p)
    details = [
        f"по направлению подготовки {metadata.program_code} {metadata.program_name}",
        f"профиль «{metadata.study_profile}»",
        "",
        f"Обучающийся: {metadata.author}, группа {metadata.student_group}",
        f"Руководитель: {metadata.supervisor_details} {metadata.supervisor_name}",
    ]
    if metadata.norm_controller:
        details.append(f"Нормоконтролер: {metadata.norm_controller}")
    elements.extend(
        _front_paragraph(document, line, style="Tex2Sto Title Details")._p
        for line in details
    )
    city = _front_paragraph(
        document,
        f"{metadata.city} {metadata.year}",
        style="Tex2Sto Title City",
    )
    city.add_run().add_break(WD_BREAK.PAGE)
    elements.append(city._p)
    return elements


def _assignment_pages(document, project: SourceProject) -> list:
    metadata = project.metadata
    if metadata.assignment_variant == "none":
        return []
    elements = []
    heading = _front_paragraph(document, "ЗАДАНИЕ", style="Tex2Sto Structural Heading")
    elements.append(heading._p)
    lines = (
        f"на выпускную квалификационную работу обучающемуся {metadata.author}",
        f"Тема: {metadata.title}",
        f"Утверждающий: {metadata.assignment_approver}",
        f"Цель работы: {metadata.assignment_goal or 'определяется содержанием работы'}",
        f"Вопросы, подлежащие разработке: {metadata.assignment_questions}",
        f"Дата выдачи: {metadata.assignment_issued}",
        f"Срок представления: {metadata.assignment_due}",
    )
    assignment_paragraphs = [
        _front_paragraph(document, line, style="Tex2Sto Assignment") for line in lines
    ]
    assignment_paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    elements.extend(paragraph._p for paragraph in assignment_paragraphs)
    if metadata.assignment_variant == "2":
        continuation = _front_paragraph(
            document,
            "ЗАДАНИЕ (ПРОДОЛЖЕНИЕ)",
            style="Tex2Sto Structural Heading",
        )
        elements.append(continuation._p)
        signature_lines = [
            "Консультанты:",
            *(f"{item.role}: {item.name}" for item in metadata.consultants),
            "",
            f"Руководитель: {metadata.supervisor_name}",
            f"Обучающийся: {metadata.author}",
        ]
        signature_paragraphs = [
            _front_paragraph(document, line, style="Tex2Sto Assignment")
            for line in signature_lines
        ]
        signature_paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        elements.extend(paragraph._p for paragraph in signature_paragraphs)
    return elements


def _abstract_and_toc(
    document,
    project: SourceProject,
    index: NumberingIndex,
    toc_entries: list[tuple[int, str, str]],
) -> list:
    counts = {kind: sum(item.kind is kind for item in index.objects) for kind in ObjectKind}
    appendices = len(find_command_calls(project.body, "appendix", 2))
    elements = []
    heading = _front_paragraph(document, "РЕФЕРАТ", style="Tex2Sto Structural Heading")
    elements.append(heading._p)
    stats = _front_paragraph(document)
    stats.add_run("Работа: ")
    add_field(stats, "NUMPAGES", "0")
    stats.add_run(
        f" с., {counts[ObjectKind.FIGURE]} рис., {counts[ObjectKind.TABLE]} табл., "
        f"{len(project.bibliography)} источник(ов), "
        f"{appendices} прил."
    )
    elements.append(stats._p)
    keywords = _front_paragraph(
        document,
        ", ".join(keyword.upper() for keyword in project.metadata.keywords),
        style="Tex2Sto Abstract Keywords",
    )
    elements.append(keywords._p)
    abstract_text = " ".join(project.metadata.abstract.split())
    abstract = _front_paragraph(document, abstract_text)
    abstract.add_run().add_break(WD_BREAK.PAGE)
    elements.append(abstract._p)
    toc_heading = _front_paragraph(document, "СОДЕРЖАНИЕ", style="Tex2Sto Structural Heading")
    elements.append(toc_heading._p)
    for level, title, bookmark in toc_entries:
        entry = _front_paragraph(document, style=f"Tex2Sto TOC {level}")
        entry.add_run(title)
        entry.add_run("\t")
        add_field(entry, f"PAGEREF {bookmark} \\h", "0")
        elements.append(entry._p)
    return elements


def _strip_paragraph_prefix(paragraph, prefix: str) -> None:
    for text in paragraph._p.iter(qn("w:t")):
        if text.text and text.text.startswith(prefix):
            text.text = text.text.removeprefix(prefix).lstrip()
            return


def _format_heading_runs(paragraph) -> None:
    for tab in list(paragraph._p.iter(qn("w:tab"))):
        tab.tag = qn("w:t")
        tab.set(qn("xml:space"), "preserve")
        tab.text = "\u00a0"
    for run in paragraph.runs:
        run.bold = False
        run.italic = False
        run_style = run._r.find(qn("w:rPr"))
        if run_style is not None:
            style = run_style.find(qn("w:rStyle"))
            if style is not None and style.get(qn("w:val")) == "SectionNumber":
                run_style.remove(style)


def _apply_styles(document) -> None:
    document.styles["Tex2Sto Table Continuation"].font.color.rgb = RGBColor(0, 0, 0)
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == PAGE_BREAK_MARKER:
            paragraph.clear()
            paragraph.add_run().add_break(WD_BREAK.PAGE)
            paragraph.style = "Tex2Sto Body"
        elif (
            text
            in {
                "ВВЕДЕНИЕ",
                "ЗАКЛЮЧЕНИЕ",
                "ОПРЕДЕЛЕНИЯ, ОБОЗНАЧЕНИЯ И СОКРАЩЕНИЯ",
                "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
            }
            or text.startswith("ПРИЛОЖЕНИЕ ")
        ):
            paragraph.style = "Tex2Sto Structural Heading"
        elif paragraph.style.name.startswith("Heading "):
            level = paragraph.style.name.rsplit(" ", 1)[-1]
            paragraph.style = f"Tex2Sto Heading {level}"
        elif paragraph._p.find(f".//{qn('w:drawing')}") is not None:
            paragraph.style = "Tex2Sto Figure Content"
        elif text.startswith(("Рисунок ", "Таблица ")):
            paragraph.style = (
                "Tex2Sto Figure Caption" if text.startswith("Рисунок ") else "Tex2Sto Table Caption"
            )
        elif text.startswith(SYMBOLS_MARKER):
            _strip_paragraph_prefix(paragraph, SYMBOLS_MARKER)
            paragraph.style = "Tex2Sto Symbols"
        elif paragraph.style.name == "Source Code":
            paragraph.style = "Tex2Sto Code"
        elif paragraph.style.name in {"Normal", "First Paragraph"}:
            paragraph.style = "Tex2Sto Body"
        if paragraph.style.name in {
            "Tex2Sto Structural Heading",
            "Tex2Sto Heading 1",
            "Tex2Sto Heading 2",
            "Tex2Sto Heading 3",
            "Tex2Sto Heading 4",
        }:
            _format_heading_runs(paragraph)
    for table in document.tables:
        table.style = "Tex2Sto Table"
        if table.rows:
            repeat_table_row(table.rows[0])
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.style = "Tex2Sto Table Text"


def _format_lists(document) -> None:
    numbering = document.part.numbering_part.element
    for abstract in numbering.findall(qn("w:abstractNum")):
        identifier = int(abstract.get(qn("w:abstractNumId"), "0"))
        if identifier < 991:
            continue
        for level in abstract.findall(qn("w:lvl")):
            number_format = level.find(qn("w:numFmt"))
            level_text = level.find(qn("w:lvlText"))
            if number_format is None or level_text is None:
                continue
            current_format = number_format.get(qn("w:val"))
            current_text = level_text.get(qn("w:val"), "")
            if current_format == "bullet":
                level_text.set(qn("w:val"), "-")
                run_properties = level.find(qn("w:rPr"))
                if run_properties is not None:
                    level.remove(run_properties)
            elif current_format == "lowerLetter":
                number_format.set(qn("w:val"), "russianLower")
            elif current_format == "lowerRoman":
                number_format.set(qn("w:val"), "decimal")
            if current_format != "bullet":
                level_text.set(qn("w:val"), current_text.rstrip(".") + ")")
            suffix = level.find(qn("w:suff"))
            if suffix is None:
                suffix = OxmlElement("w:suff")
                level.append(suffix)
            suffix.set(qn("w:val"), "nothing")
    for paragraph in document.paragraphs:
        properties = paragraph._p.pPr
        if properties is None or properties.find(qn("w:numPr")) is None:
            continue
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = "\u00a0"
        run.append(text)
        paragraph._p.insert(1, run)


def _toc_entries(document) -> list[tuple[int, str, str]]:
    headings: list[tuple[int, str, object]] = []
    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        style = paragraph.style.name
        if not text:
            continue
        if style == "Tex2Sto Structural Heading":
            headings.append((1, text, paragraph))
        elif style.startswith("Tex2Sto Heading "):
            headings.append((int(style.rsplit(" ", 1)[-1]), text, paragraph))

    entries: list[tuple[int, str, str]] = []
    for position, (level, title, paragraph) in enumerate(headings):
        if title.startswith("ПРИЛОЖЕНИЕ "):
            parts = title.split()
            title = f"Приложение {parts[1]}"
        else:
            title = STRUCTURAL_TOC_TITLES.get(title, title)
        bookmark = f"tex2sto_toc_{position + 1}"
        add_bookmark(paragraph, bookmark, position + 100)
        entries.append((level, title, bookmark))
    return entries


def _format_equations(document, index: NumberingIndex) -> None:
    numbers = [item.display for item in index.objects if item.kind is ObjectKind.EQUATION]
    equation_index = 0
    for paragraph in document.paragraphs:
        math_paragraph = paragraph._p.find(qn("m:oMathPara"))
        if math_paragraph is None:
            continue
        math = math_paragraph.find(qn("m:oMath"))
        if math is None or equation_index >= len(numbers):
            continue
        math_paragraph.remove(math)
        table = document.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        widths = (Mm(30), Mm(105), Mm(30))
        table_width = table._tbl.tblPr.find(qn("w:tblW"))
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(sum(round(width.twips) for width in widths)))
        grid_columns = table._tbl.tblGrid.findall(qn("w:gridCol"))
        for position, (cell, width) in enumerate(zip(table.rows[0].cells, widths, strict=True)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            grid_columns[position].set(qn("w:w"), str(round(width.twips)))
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "0")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
        center = table.cell(0, 1).paragraphs[0]
        center.style = "Tex2Sto Equation"
        center.alignment = WD_ALIGN_PARAGRAPH.CENTER
        center._p.append(math)
        number = table.cell(0, 2).paragraphs[0]
        number.style = "Tex2Sto Equation"
        number.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        number.add_run(f"({numbers[equation_index]})")
        parent = paragraph._p.getparent()
        insertion = parent.index(paragraph._p)
        table._tbl.getparent().remove(table._tbl)
        parent.insert(insertion, table._tbl)
        parent.remove(paragraph._p)
        equation_index += 1


def _balance_table_columns(table) -> None:
    grid_columns = table._tbl.tblGrid.findall(qn("w:gridCol"))
    if len(grid_columns) < 2:
        return
    total_width = sum(int(column.get(qn("w:w"), "0")) for column in grid_columns)
    if total_width <= 0:
        return

    content_rows = [
        row
        for row in table.rows
        if not any(TABLE_BREAK_MARKER in cell.text for cell in row.cells)
    ]
    weights = [
        max((len(row.cells[position].text.strip()) for row in content_rows), default=0) + 8
        for position in range(len(grid_columns))
    ]
    allocated = 0
    for position, (column, weight) in enumerate(zip(grid_columns, weights, strict=True)):
        width = (
            total_width - allocated
            if position == len(grid_columns) - 1
            else round(total_width * weight / sum(weights))
        )
        column.set(qn("w:w"), str(width))
        allocated += width


def _table_segment(table, start: int, end: int):
    segment = deepcopy(table._tbl)
    for position, row in enumerate(segment.findall(qn("w:tr"))):
        if position != 0 and not start <= position < end:
            segment.remove(row)
    return segment


def _split_longtable(document, table, display: str) -> None:
    marker_rows = [
        position
        for position, row in enumerate(table.rows)
        if any(TABLE_BREAK_MARKER in cell.text for cell in row.cells)
    ]
    if not marker_rows:
        raise RuntimeError("validated longtable lost its table-break marker during DOCX rendering")

    starts = [1, *(position + 1 for position in marker_rows)]
    ends = [*marker_rows, len(table.rows)]
    elements = []
    for segment_number, (start, end) in enumerate(zip(starts, ends, strict=True)):
        if segment_number:
            continuation = _front_paragraph(document, style="Tex2Sto Table Continuation")
            continuation.add_run().add_break(WD_BREAK.PAGE)
            continuation.add_run(f"Продолжение таблицы {display}")
            elements.append(continuation._p)
        elements.append(_table_segment(table, start, end))

    parent = table._tbl.getparent()
    insertion = parent.index(table._tbl)
    parent.remove(table._tbl)
    for element in elements:
        parent.insert(insertion, element)
        insertion += 1


def _format_tables(document, project: SourceProject, index: NumberingIndex) -> None:
    table_objects = [item for item in index.objects if item.kind is ObjectKind.TABLE]
    longtable_positions = {start for start, _, _ in find_environment(project.body, "longtable")}
    captions = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("Таблица ")
    ]
    tables = list(document.tables)
    for position, (table, numbered) in enumerate(zip(tables, table_objects, strict=False)):
        _balance_table_columns(table)
        bookmark = f"tex2sto_table_{position + 1}"
        if position < len(captions):
            add_bookmark(captions[position], bookmark, position + 1000)
        if numbered.position not in longtable_positions or not table.rows:
            continue
        _split_longtable(document, table, numbered.display)


def _postprocess(path: Path, project: SourceProject, index: NumberingIndex) -> None:
    document = Document(path)
    _apply_styles(document)
    _format_lists(document)
    _format_tables(document, project, index)
    _format_equations(document, index)
    toc_entries = _toc_entries(document)
    anchor = document.element.body[0]
    elements = [
        *_title_page(document, project),
        *_assignment_pages(document, project),
        *_abstract_and_toc(document, project, index, toc_entries),
    ]
    for element in elements:
        _move_before(element, anchor)
    set_update_fields(document)
    document.core_properties.title = project.metadata.title
    document.core_properties.author = project.metadata.author
    document.save(path)


def render_docx(project: SourceProject, index: NumberingIndex, output: Path) -> Path:
    pandoc = _pandoc()
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="tex2sto-docx-") as temporary:
        source = Path(temporary) / "document.tex"
        source.write_text(renderer_body(project, index, target="docx"), encoding="utf-8")
        command = [
            pandoc,
            str(source),
            "--from=latex",
            "--to=docx",
            "--standalone",
            "--number-sections",
            f"--reference-doc={PROFILE / 'reference.docx'}",
            f"--lua-filter={PROFILE / 'ssau.lua'}",
            f"--resource-path={project.root}",
            "--output",
            str(output),
        ]
        run_tool(command, cwd=project.root)
    _postprocess(output, project, index)
    return output
